#!/usr/bin/env python3
"""Convert book Markdown to a styled Word file for Google Docs."""

from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

NAVY = RGBColor(0x1F, 0x5C, 0x40)
INK = RGBColor(0x24, 0x30, 0x28)
MUTED = RGBColor(0x5F, 0x70, 0x68)
ACCENT = "EEF8F2"
HEADER_FILL = "3D8F62"
HEADER_TEXT = RGBColor(0xFF, 0xFF, 0xFF)

BOLD = re.compile(r"\*\*(.+?)\*\*")
MARKDOWN_LINK = re.compile(r"\[([^\]]+)\]\((https?://[^)]+)\)")
ANGLE_LINK = re.compile(r"<(https?://[^>]+)>")
BILD_RE = re.compile(
    r"^(?:\*\*)?Bild:(?:\*\*)?\s*`([^`]+)`(?:\s*[—–-]\s*(.*))?",
    re.I,
)


def set_font(run, name="Calibri") -> None:
    run.font.name = name
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set(qn("w:ascii"), name)
    rFonts.set(qn("w:hAnsi"), name)
    rFonts.set(qn("w:eastAsia"), name)


def shade_paragraph(paragraph, fill: str) -> None:
    pPr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), fill)
    pPr.append(shd)


def shade_cell(cell, fill: str) -> None:
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)


def add_runs(paragraph, text: str, color=INK, size=Pt(12)) -> None:
    text = MARKDOWN_LINK.sub(r"\1 — \2", text)
    text = ANGLE_LINK.sub(r"\1", text)
    parts = BOLD.split(text)
    for i, part in enumerate(parts):
        if not part:
            continue
        run = paragraph.add_run(part)
        run.bold = i % 2 == 1
        run.font.size = size
        run.font.color.rgb = color
        set_font(run)


def style_heading(style, size, color, space_before, space_after) -> None:
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = color
    style.font.name = "Calibri"
    style.paragraph_format.space_before = Pt(space_before)
    style.paragraph_format.space_after = Pt(space_after)
    style.paragraph_format.line_spacing = 1.15


def set_document_style(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.3)
    section.bottom_margin = Cm(2.3)
    section.left_margin = Cm(2.4)
    section.right_margin = Cm(2.4)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(12)
    normal.font.color.rgb = INK
    normal.paragraph_format.space_after = Pt(10)
    normal.paragraph_format.line_spacing = 1.15
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE

    style_heading(doc.styles["Heading 1"], 26, NAVY, 6, 14)
    style_heading(doc.styles["Heading 2"], 16, NAVY, 18, 8)
    style_heading(doc.styles["Heading 3"], 13, NAVY, 14, 6)


def set_left_border(paragraph, color="2D6A4F", sz="24") -> None:
    pPr = paragraph._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), sz)
    left.set(qn("w:space"), "10")
    left.set(qn("w:color"), color)
    pBdr.append(left)
    pPr.append(pBdr)


def add_cover(doc: Document) -> None:
    table = doc.add_table(rows=1, cols=1)
    cell = table.cell(0, 0)
    shade_cell(cell, "F4FAF6")
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(36)
    run = p.add_run("INSTRUKTÖRSVERKET AB  ·  TÄNNDALEN")
    set_font(run)
    run.font.size = Pt(11)
    run.font.color.rgb = RGBColor(0x3D, 0x8F, 0x62)
    run.bold = True

    p2 = cell.add_paragraph()
    p2.paragraph_format.space_before = Pt(14)
    r2 = p2.add_run("Förarbevis\nsnöskoter")
    set_font(r2)
    r2.font.size = Pt(36)
    r2.bold = True
    r2.font.color.rgb = RGBColor(0x1F, 0x5C, 0x40)

    p3 = cell.add_paragraph()
    p3.paragraph_format.space_before = Pt(10)
    p3.paragraph_format.space_after = Pt(14)
    r3 = p3.add_run("Utbildningsmaterial")
    set_font(r3)
    r3.font.size = Pt(13)
    r3.font.color.rgb = RGBColor(0x4A, 0x6B, 0x5A)

    cover_img = Path(__file__).resolve().parents[1] / "assets" / "omslag-skoter.jpg"
    if cover_img.exists():
        p_img = cell.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(8)
        p_img.paragraph_format.space_after = Pt(14)
        p_img.add_run().add_picture(str(cover_img), width=Cm(16))

    p4 = cell.add_paragraph()
    r4 = p4.add_run("Teori och praktik  ·  Förarbevis snöskoter")
    set_font(r4)
    r4.font.size = Pt(11)
    r4.font.color.rgb = RGBColor(0x4A, 0x6B, 0x5A)
    doc.add_page_break()


def add_page_numbers(doc: Document) -> None:
    footer = doc.sections[0].footer
    footer.is_linked_to_previous = False
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("Instruktörsverket AB  ·  Förarbevis snöskoter  ·  ")
    set_font(run)
    run.font.size = Pt(9)
    run.font.color.rgb = MUTED
    fld1 = OxmlElement("w:fldChar")
    fld1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld2 = OxmlElement("w:fldChar")
    fld2.set(qn("w:fldCharType"), "end")
    r2 = p.add_run()
    r2._r.append(fld1)
    r2._r.append(instr)
    r2._r.append(fld2)
    set_font(r2)
    r2.font.size = Pt(9)
    r2.font.color.rgb = MUTED


def parse_row(line: str) -> list[str]:
    return [cell.strip() for cell in line.strip("|").split("|")]


def add_markdown(doc: Document, lines: list[str], assets_dir: Path | None = None) -> None:
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        if not line:
            i += 1
            continue

        if line.startswith("|") and line.endswith("|"):
            rows = []
            while i < len(lines) and lines[i].rstrip().startswith("|"):
                raw = lines[i].rstrip()
                if not re.match(r"^\|[\s:|-]+\|$", raw):
                    rows.append(parse_row(raw))
                i += 1
            if rows:
                table = doc.add_table(rows=len(rows), cols=len(rows[0]))
                table.style = "Table Grid"
                for r, row in enumerate(rows):
                    for c, value in enumerate(row):
                        cell = table.rows[r].cells[c]
                        cell.text = ""
                        p = cell.paragraphs[0]
                        if r == 0:
                            shade_cell(cell, HEADER_FILL)
                            add_runs(p, value, color=HEADER_TEXT, size=Pt(11))
                            for run in p.runs:
                                run.bold = True
                        else:
                            if r % 2 == 0:
                                shade_cell(cell, "F4F7F4")
                            add_runs(p, value, size=Pt(11))
                doc.add_paragraph()
            continue

        if line.startswith("# "):
            p = doc.add_heading(line[2:].strip(), level=1)
            shade_paragraph(p, "D8EEE3")
            for run in p.runs:
                set_font(run)
                run.font.color.rgb = RGBColor(0x1F, 0x5C, 0x40)
                run.font.size = Pt(22)
        elif line.startswith("## "):
            p = doc.add_heading(line[3:].strip(), level=2)
            for run in p.runs:
                set_font(run)
                run.font.color.rgb = NAVY
        elif line.startswith("### "):
            p = doc.add_heading(line[4:].strip(), level=3)
            for run in p.runs:
                set_font(run)
                run.font.color.rgb = RGBColor(0x2D, 0x6A, 0x4F)
        elif line.startswith("> "):
            body = line[2:].strip()
            m = BILD_RE.match(body)
            img_path = None
            caption = ""
            if m and assets_dir is not None:
                img_path = assets_dir / Path(m.group(1).strip()).name
                caption = (m.group(2) or "").strip()
                if not img_path.exists():
                    img_path = None
            if img_path is not None:
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_before = Pt(8)
                p.paragraph_format.space_after = Pt(4)
                run = p.add_run()
                width = Cm(6.2) if img_path.name.startswith(("marke-", "ledkryss", "vagmarken")) else Cm(14.2)
                run.add_picture(str(img_path), width=width)
                if caption:
                    cap = doc.add_paragraph()
                    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    cap.paragraph_format.space_after = Pt(12)
                    add_runs(cap, caption, color=MUTED, size=Pt(10))
            else:
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Cm(0.15)
                p.paragraph_format.space_before = Pt(8)
                p.paragraph_format.space_after = Pt(12)
                remember = "kom ihåg" in body.lower()
                shade_paragraph(p, "F8EFE8" if remember else ACCENT)
                set_left_border(p, "C45C26" if remember else "2D6A4F")
                add_runs(p, body, color=NAVY)
                for run in p.runs:
                    run.italic = True
        elif line.startswith("- "):
            body = line[2:].strip()
            i += 1
            while i < len(lines) and lines[i].startswith(("  ", "\t")):
                body += " " + lines[i].strip()
                i += 1
            p = doc.add_paragraph(style="List Bullet")
            add_runs(p, body)
            continue
        elif re.match(r"^\d+\. ", line):
            body = re.sub(r"^\d+\. ", "", line).strip()
            i += 1
            while i < len(lines) and lines[i].startswith(("  ", "\t")):
                body += " " + lines[i].strip()
                i += 1
            p = doc.add_paragraph(style="List Number")
            add_runs(p, body)
            continue
        else:
            p = doc.add_paragraph()
            add_runs(p, line)

        i += 1


def convert(md_path: Path, docx_path: Path) -> None:
    doc = Document()
    set_document_style(doc)
    add_markdown(doc, md_path.read_text(encoding="utf-8").splitlines(), md_path.parents[1] / "assets" if md_path.parent.name == "kapitel" else md_path.parent / "assets")
    docx_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(docx_path)


def convert_many(md_paths: list[Path], docx_path: Path) -> None:
    doc = Document()
    set_document_style(doc)
    add_cover(doc)
    assets = md_paths[0].parent.parent / "assets"
    for i, path in enumerate(md_paths):
        if i:
            doc.add_page_break()
        add_markdown(doc, path.read_text(encoding="utf-8").splitlines(), assets)
    add_page_numbers(doc)
    docx_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(docx_path)


BOOK_FILES = [
    "00-forord.md",
    "01-forare-och-fordon.md",
    "02-miljo-och-natur.md",
    "03-trafikregler.md",
    "04-straff-behorighet-forsakring.md",
    "05-trafiksakerhet.md",
    "06-manniskan.md",
    "07-olyckor-forsta-hjalpen.md",
    "08-praktik.md",
    "09-infor-provet.md",
]


if __name__ == "__main__":
    root = Path(__file__).resolve().parents[1]
    kapitel = root / "kapitel"
    if len(sys.argv) > 1 and sys.argv[1] == "--bok":
        out = root / "export" / "Forarbevis-snoskoter-Instruktorsverket.docx"
        convert_many([kapitel / name for name in BOOK_FILES], out)
    elif len(sys.argv) > 1:
        src = Path(sys.argv[1])
        out = (
            Path(sys.argv[2])
            if len(sys.argv) > 2
            else root / "export" / (src.stem + ".docx")
        )
        convert(src, out)
    else:
        out = root / "export" / "Forarbevis-snoskoter-Instruktorsverket.docx"
        convert_many([kapitel / name for name in BOOK_FILES], out)
    print(out)
